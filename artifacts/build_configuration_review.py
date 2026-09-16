from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "artifacts/Evaluation Configuration Bug and Change List.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_font(run, size=10.0, bold=False, color="000000"):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_cell(cell, text, *, bold=False, color="000000", size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.04
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


def clear_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def clear_style_border(style):
    p_pr = style._element.find(qn("w:pPr"))
    if p_pr is None:
        return
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size=14 if level == 1 else 11.5, bold=True)
    return p


def add_body(doc, text, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_font(r, size=10.4)
    return p


def add_bullets(doc, items):
    for text in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(text)
        set_font(r, size=10.2)


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=5)
    table.autofit = False
    table.style = "Table Grid"
    table.allow_autofit = False
    widths = [Inches(0.62), Inches(0.82), Inches(0.78), Inches(2.63), Inches(2.35)]
    header_cells = table.rows[0].cells
    headers = ["ID", "Area", "Priority", "Finding and risk", "Recommended change"]
    for cell, label, width in zip(header_cells, headers, widths):
        cell.width = width
        set_cell_shading(cell, "1F4E78")
        style_cell(cell, label, bold=True, color="FFFFFF", size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_header(table.rows[0])
    for index, row_data in enumerate(rows):
        cells = table.add_row().cells
        fill = "F4F8FB" if index % 2 else "FFFFFF"
        for cell, value, width in zip(cells, row_data, widths):
            cell.width = width
            set_cell_shading(cell, fill)
            style_cell(cell, value)
        style_cell(cells[0], row_data[0], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        priority_color = {"Critical": "B91C1C", "High": "A14F00", "Medium": "1D4ED8"}[row_data[2]]
        style_cell(cells[2], row_data[2], bold=True, color=priority_color, align=WD_ALIGN_PARAGRAPH.CENTER)
        keep_row_together(table.rows[-1])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.4)
    clear_style_border(doc.styles["Title"])

    title = doc.add_paragraph(style="Title")
    clear_paragraph_border(title)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(7)
    run = title.add_run("Evaluation Configuration Bug and Change List")
    set_font(run, size=21, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Academic years Regulations Terms Exam sessions Evaluation events Programmes Courses Subjects Evaluation centres Papers")
    set_font(run, size=10.5, color="404040")

    add_body(
        doc,
        "Fix the critical paper controls and event-centre validation before using this configuration module for a live examination. The current implementation can allow unreviewed paper content or an event with no assigned centre to appear ready for go-live.",
        after=8,
    )

    add_heading(doc, "Review basis")
    add_body(doc, "The review covered the configuration workspace, configuration service layer, API routes, models, and automated workflow tests. The existing configuration test suite completed successfully with 5 passing tests.")

    add_heading(doc, "Immediate changes")
    add_bullets(doc, [
        "Prevent all question and paper-content changes once a paper leaves Draft. Route post-submission changes through the governed change process.",
        "Require a different administrator to approve and freeze a paper than the person who submitted it.",
        "Require every active evaluation event to have at least one active assigned centre. Include that relationship in the go-live check.",
        "Add controlled edit and retire actions for every master record, not only for papers.",
    ])

    add_heading(doc, "Detailed findings")
    findings = [
        ("CFG 01", "Papers", "Critical", "A question can be added while a paper is In review or Approved because the service blocks changes only when the paper is Frozen. The changed paper does not return to review.", "Allow question changes only in Draft. For later changes, create a governed revision and require fresh review and approval."),
        ("CFG 02", "Papers", "Critical", "One administrator can create, submit, approve, and freeze the same paper. The passing workflow test uses one authenticated client for the whole lifecycle, confirming no separation check exists.", "Store the submitting actor. Reject self-approval and require a different actor for approval and freeze, with role checks and audit evidence."),
        ("CFG 03", "Events Centres", "Critical", "An evaluation event accepts an empty centre list. Go-live readiness checks only that any active event and any active centre exist, not that the event has an assigned usable centre.", "Require one or more unique active centres for every active event. Validate every assigned centre again in readiness."),
        ("CFG 04", "Master data", "High", "Academic years, regulations, terms, sessions, events, programmes, courses, subjects, and centres are create-only in both the workspace and API. Data-entry mistakes cannot be corrected or retired.", "Add governed edit, activate, deactivate, and retire actions with dependency checks, version history, and audit events for every master entity."),
        ("CFG 05", "Exam sessions", "High", "The session term is free text instead of a link to a configured Term. Session dates only need to be ordered; they can sit outside the selected academic year or conflict with the named term.", "Replace the text field with a Term foreign key. Enforce that the term belongs to the selected academic year and contains the session dates."),
        ("CFG 06", "Session status", "High", "The model has Draft, Approval, Ready, Active, and Closed statuses, but the configuration API and workspace have no status transition. Readiness can fall back to the latest Draft session.", "Implement a session lifecycle with explicit readiness approval and activation. Do not return Go live for a Draft, Approval, or Closed session."),
        ("CFG 07", "Subjects Papers", "High", "A paper can be created for any tenant subject and any tenant session. The subject Available exam sessions list is never checked when a paper is mapped.", "Reject a paper when its session is not available for that subject. Define a clear rule for subjects with no listed sessions."),
        ("CFG 08", "Papers", "High", "Paper creation bypasses the positive-maximum validation used by later changes. A zero-mark paper with zero-mark questions can be marked ready and progress through the lifecycle.", "Run the full snapshot validation during creation. Require maximum marks above zero and enforce all numeric ranges before saving a draft."),
        ("CFG 09", "Programmes Courses", "High", "Programme regulation is a free-text value, while Course regulation is a linked record. This permits a programme and its courses to carry conflicting regulations.", "Change Programme regulation to a Regulation foreign key and require each course to use the programme regulation unless an approved exception exists."),
        ("CFG 10", "Courses", "Medium", "Course duration accepts zero because both the client and server accept non-negative values. Courses can also use inactive or expired regulations without a validation check.", "Require duration terms to be at least one. Validate that the linked regulation is active and effective for the course configuration."),
        ("CFG 11", "Subjects", "Medium", "Related subjects can include the same subject or subjects from unrelated programmes or courses. The service verifies tenant membership only.", "Reject self-links and require an explicit compatibility rule for cross-programme or cross-course relationships."),
        ("CFG 12", "Academic years Terms", "Medium", "Academic years can overlap, and terms only need to be inside a year. Overlapping terms and inconsistent academic calendars are accepted.", "Add date-overlap validation for active years and terms, with an approved exception workflow for deliberate overlaps."),
        ("CFG 13", "Evaluation centres", "Medium", "Approved network CIDRs are stored as unvalidated text and are not used elsewhere in the codebase. Invalid entries can be saved and centre network restrictions are not enforced.", "Validate each CIDR with a network parser, show field-level errors, and apply the approved-network rule to centre-bound operational access."),
        ("CFG 14", "Configuration view", "High", "The workspace hides essential relationships: terms do not show their academic year, events do not show session or centres, and papers do not show session or subject. The readiness banner shows only a count, not the problems to fix.", "Show human-readable parent records and assigned centres in each list. Expand the readiness panel to display every blocking issue with a direct link to the relevant record."),
        ("CFG 15", "Sessions Events", "Medium", "Sessions and evaluation events have no uniqueness or idempotency protection. A retry after a slow network response can create duplicate operational records.", "Add natural unique constraints and idempotency keys for create requests, then return the existing result safely on a retry."),
    ]
    add_table(doc, findings)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("Configuration review 16 September 2026")
    set_font(r, size=8.5, color="666666")

    doc.core_properties.title = "Evaluation Configuration Bug and Change List"
    doc.core_properties.subject = "Prioritized configuration issues and recommended changes"
    doc.core_properties.author = "ADMIEZO configuration review"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
